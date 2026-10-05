import io
import json

import pytest
from django.contrib.auth.models import Permission
from django.contrib.messages import get_messages
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from encuestas.models import (
    Encuesta,
    CumplimientoRonda,
    EstadoEncuesta,
    Pregunta,
    TipoPregunta,
    TipoSegmentacion,
)
from encuestas.services import (
    actualizar_segmentacion,
    agregar_destinatario,
    crear_encuesta,
    exportar_encuesta,
    get_rondas_pendientes,
    importar_encuesta,
    nueva_version,
    quitar_destinatario,
    usuario_esta_segmentado,
)
from encuestas.tests.helpers import publicar_para_test
from encuestas.validators import generar_plantilla_listado

pytestmark = pytest.mark.django_db
TIPO = TipoSegmentacion.LISTADO_USUARIOS


@pytest.fixture
def usuarios(django_user_model):
    return [
        django_user_model.objects.create_user(username=f"id-test-{i}") for i in range(3)
    ]


@pytest.fixture
def encuesta(usuarios):
    obj = crear_encuesta(
        usuario=usuarios[0],
        titulo="Por IDs",
        es_obligatoria=True,
        duracion_ronda_dias=7,
    )
    Pregunta.objects.create(
        encuesta=obj, texto="¿Todo bien?", tipo=TipoPregunta.SI_NO, orden=1
    )
    return obj


def csv(contenido):
    return SimpleUploadedFile("usuarios.csv", contenido.encode("utf-8"))


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_carga_deduplica_y_reemplaza(encuesta, usuarios, extension):
    actualizar_segmentacion(encuesta, tipo=TIPO, usuarios_ids=[usuarios[0].pk])
    if extension == "csv":
        archivo = csv(
            f"usuario_id\n{usuarios[1].pk}\n\n{usuarios[1].pk}\n{usuarios[2].pk}\n"
        )
    else:
        book = Workbook()
        for row in [
            ["usuario_id"],
            [usuarios[1].pk],
            [],
            [usuarios[1].pk],
            [usuarios[2].pk],
        ]:
            book.active.append(row)
        stream = io.BytesIO()
        book.save(stream)
        archivo = SimpleUploadedFile("usuarios.xlsx", stream.getvalue())
    seg = actualizar_segmentacion(encuesta, tipo=TIPO, archivo=archivo)
    assert set(seg.usuarios.values_list("pk", flat=True)) == {
        u.pk for u in usuarios[1:]
    }


@pytest.mark.parametrize(
    "contenido",
    [
        "usuario_id\n0",
        "usuario_id\n-1",
        "usuario_id\n1.5",
        "usuario_id\nabc",
        "usuario_id\n999999999999999999999999",
        "usuario_id\n99999999",
        "numero_documento\n123",
        "usuario_id\n",
        "usuario_id,otra\n,x",
    ],
)
def test_archivo_invalido_no_cambia_segmentacion(encuesta, usuarios, contenido):
    actualizar_segmentacion(
        encuesta,
        tipo=TipoSegmentacion.LISTADO_DOCUMENTOS,
        destinatarios=[{"tipo_documento": "dni", "numero_documento": "123"}],
    )
    with pytest.raises(ValidationError):
        actualizar_segmentacion(encuesta, tipo=TIPO, archivo=csv(contenido))
    encuesta.refresh_from_db()
    assert encuesta.segmentacion.tipo == TipoSegmentacion.LISTADO_DOCUMENTOS
    assert encuesta.segmentacion.destinatarios.count() == 1
    assert not encuesta.segmentacion.usuarios.exists()


def test_manual_y_cambio_de_tipo_limpian_destinatarios(encuesta, usuarios):
    actualizar_segmentacion(encuesta, tipo=TIPO)
    for _ in range(2):
        agregar_destinatario(encuesta, usuario_id=str(usuarios[1].pk))
    assert encuesta.segmentacion.usuarios.count() == 1
    with pytest.raises(ValidationError):
        agregar_destinatario(encuesta, usuario_id="99999999")
    with pytest.raises(ValidationError):
        quitar_destinatario(encuesta, usuarios[2].pk)
    quitar_destinatario(encuesta, usuarios[1].pk)
    assert not encuesta.segmentacion.usuarios.exists()
    agregar_destinatario(encuesta, usuario_id=usuarios[1].pk)
    actualizar_segmentacion(encuesta, tipo=TipoSegmentacion.TODOS_LOS_USUARIOS)
    assert not encuesta.segmentacion.usuarios.exists()
    actualizar_segmentacion(encuesta, tipo=TIPO)
    assert not usuario_esta_segmentado(encuesta, usuarios[1])


def test_pendientes_y_respuesta_solo_para_usuario_incluido(client, encuesta, usuarios):
    actualizar_segmentacion(encuesta, tipo=TIPO, usuarios_ids=[usuarios[1].pk])
    ronda = publicar_para_test(encuesta, usuario=usuarios[0])
    assert usuario_esta_segmentado(encuesta, usuarios[1])
    assert not usuario_esta_segmentado(encuesta, usuarios[2])
    assert [r.pk for r in get_rondas_pendientes(usuarios[1])] == [ronda.pk]
    assert get_rondas_pendientes(usuarios[2]) == []
    client.force_login(usuarios[1])
    url = reverse("encuestas_responder", args=[ronda.pk])
    datos = {f"respuesta-{encuesta.preguntas.get().pk}": "si"}
    assert client.post(url, datos).status_code == 302
    assert get_rondas_pendientes(usuarios[1]) == []
    client.force_login(usuarios[2])
    response = client.post(url, datos)
    assert response.status_code == 302
    assert any(
        "No estás habilitado" in str(m) for m in get_messages(response.wsgi_request)
    )
    assert not CumplimientoRonda.objects.filter(
        ronda=ronda, usuario=usuarios[2]
    ).exists()


def test_bloqueo_pendiente(encuesta, usuarios):
    actualizar_segmentacion(encuesta, tipo=TIPO, usuarios_ids=[usuarios[1].pk])
    Encuesta.objects.filter(pk=encuesta.pk).update(
        estado=EstadoEncuesta.PENDIENTE_APROBACION
    )
    for accion in [
        lambda: actualizar_segmentacion(encuesta, tipo=TIPO, usuarios_ids=[]),
        lambda: agregar_destinatario(encuesta, usuario_id=usuarios[2].pk),
        lambda: quitar_destinatario(encuesta, usuarios[1].pk),
    ]:
        with pytest.raises(ValidationError):
            accion()
    assert encuesta.segmentacion.usuarios.count() == 1


def test_clonacion_y_portabilidad(encuesta, usuarios):
    actualizar_segmentacion(encuesta, tipo=TIPO, usuarios_ids=[usuarios[1].pk])
    nueva = nueva_version(encuesta, usuario=usuarios[0])
    assert list(nueva.segmentacion.usuarios.all()) == [usuarios[1]]
    datos = exportar_encuesta(encuesta, incluir_segmentacion=True)
    assert datos["version_formato"] == 4
    archivo = SimpleUploadedFile("encuesta.json", json.dumps(datos).encode())
    copia = importar_encuesta(archivo, usuario=usuarios[0])
    assert list(copia.segmentacion.usuarios.all()) == [usuarios[1]]
    datos["segmentacion"]["destinatarios"] = [{"usuario_id": 99999999}]
    cantidad = Encuesta.objects.count()
    with pytest.raises(ValidationError):
        importar_encuesta(
            SimpleUploadedFile("encuesta.json", json.dumps(datos).encode()),
            usuario=usuarios[0],
        )
    assert Encuesta.objects.count() == cantidad


def test_vistas_plantilla_carga_revision_y_permisos(client, encuesta, usuarios):
    plantilla = (
        reverse("encuestas_segmentacion_plantilla", args=[encuesta.pk])
        + "?tipo=listado_usuarios"
    )
    client.force_login(usuarios[0])
    assert client.get(plantilla).status_code == 403
    usuarios[0].user_permissions.add(
        *Permission.objects.filter(
            content_type__app_label="encuestas",
            codename__in=["change_encuesta", "view_encuesta"],
        )
    )
    response = client.get(plantilla)
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert list(book.active.values) == [("usuario_id",)]
    book.close()
    client.post(
        reverse("encuestas_segmentacion_tipo", args=[encuesta.pk]), {"tipo": TIPO}
    )
    client.post(
        reverse("encuestas_segmentacion_agregar", args=[encuesta.pk]),
        {"usuario_id": usuarios[1].pk},
    )
    response = client.get(reverse("encuestas_segmentacion", args=[encuesta.pk]))
    assert usuarios[1].username in response.content.decode()
    response = client.get(reverse("encuestas_revision", args=[encuesta.pk]))
    assert usuarios[1].username in response.content.decode()
    client.post(
        reverse("encuestas_segmentacion_quitar", args=[encuesta.pk, usuarios[1].pk])
    )
    encuesta.refresh_from_db()
    assert not encuesta.segmentacion.usuarios.exists()


def test_plantilla_documentos_conserva_formato():
    book = load_workbook(io.BytesIO(generar_plantilla_listado()))
    assert list(book.active.values)[0] == ("tipo_documento", "numero_documento")
    book.close()
