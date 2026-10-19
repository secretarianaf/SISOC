"""Deriva la plantilla docxtpl del Seguimiento SIPH a partir del docx del área (#2651).

Reemplaza los blancos punteados, las casillas y las celdas vacías del
"IV-Formulario Seguimiento de SIPH-Línea Tradicional y Módulos secos" por tags
Jinja (las claves que arma ``relevamientos.siph_docx.contexto_docx_siph``) y
quita las fuentes embebidas (9 MB que Word sustituye sin problema).

Se corre una sola vez cuando el área cambia el formulario; el resultado va a
``relevamientos/templates/relevamientos/docx/seguimiento_siph.docx``.

Uso (dentro del contenedor de Django, que tiene python-docx):

    python -I src/backends/comedores/relevamientos/scripts/armar_plantilla_siph.py \
        <docx_del_area> src/backends/comedores/relevamientos/templates/relevamientos/docx/seguimiento_siph.docx
"""

import re
import sys
import tempfile
import zipfile

from docx import Document
from docx.oxml.ns import qn

PUNTOS = re.compile(r"[….]{4,}")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

# (patrón sobre el texto del párrafo, texto nuevo). Se aplica el primero que
# matchea; el sangrado original del párrafo se conserva.
REGLAS = [
    (
        r"^\s*Nombre de la Organización Solicitante:",
        "Nombre de la Organización Solicitante: {{ organizacion }}",
    ),
    (r"^\s*Código Proyecto:", "Código Proyecto: {{ codigo_proyecto }}"),
    (
        r"^\s*Nombre del Punto de Entrega/Espacio Comunitario:",
        "Nombre del Punto de Entrega/Espacio Comunitario: {{ comedor_nombre }}",
    ),
    (
        r"^\s*Código del Punto de Entrega/Espacio Comunitario \(ID Externo\):",
        "Código del Punto de Entrega/Espacio Comunitario (ID Externo): "
        "{{ codigo_externo }}",
    ),
    (r"^\s*ID SISOC:", "ID SISOC: {{ id_sisoc }}"),
    (r"^\s*ID Relevamiento:", "ID Relevamiento: {{ id_relevamiento }}"),
    (
        r"^\s*Fecha y horario de la visita",
        "Fecha y horario de la visita (Día/ mes/ año) (hh:mm): {{ fecha_hora }}",
    ),
    (r"^Abierto, con actividad", "{{ x_abierto_con }} Abierto, con actividad"),
    (r"^Abierto, sin actividad", "{{ x_abierto_sin }} Abierto, sin actividad"),
    (r"^Cerrado", "{{ x_cerrado }} Cerrado"),
    (r"^Motivo: ", "Motivo: {{ motivo }}"),
    (
        r"^¿La actividad visitada se realiza",
        "¿La actividad visitada se realiza en el espacio comunitario?  "
        "SÍ {{ cb_actividad_si }} NO {{ cb_actividad_no }}",
    ),
    (
        r"^En caso de responder NO, especificar el motivo",
        "En caso de responder NO, especificar el motivo: "
        "{{ actividad_en_espacio_motivo }}",
    ),
    (r"^Nombre de la actividad:", "Nombre de la actividad: {{ actividad_nombre }}"),
    (
        r"^¿El financiamiento recibido fue utilizado",
        "¿El financiamiento recibido fue utilizado según lo planificado en los "
        "últimos tres meses?:  Sí {{ cb_fin_si }} No {{ cb_fin_no }}",
    ),
    (
        r"^◻ Retribución talleristas",
        "{{ cb_comp_talleristas }} Retribución talleristas. Detallar cuántos: "
        "{{ financiamiento_talleristas_cantidad }}",
    ),
    (
        r"^\s*◻ Materiales/insumos/equipamiento",
        "{{ cb_comp_materiales }} Materiales/insumos/equipamiento. Detallar cuáles: "
        "{{ financiamiento_materiales_detalle }}",
    ),
    (
        r"^\s*◻ Adecuación mínima",
        "{{ cb_comp_adecuacion }} Adecuación mínima de servicios básicos. "
        "Detallar cuáles: {{ financiamiento_adecuacion_detalle }}",
    ),
    (
        r"^5\.2 En caso de responder NO, justificar",
        "5.2 En caso de responder NO, justificar: {{ financiamiento_justificacion }}",
    ),
    (
        r"^◻ Si\s+◻ No",
        "{{ cb_otros_si }} Si   {{ cb_otros_no }} No\t(En caso afirmativo pasar a 5.3.1)",
    ),
    (r"^◻ Estado Nacional", "{{ cb_origen_nacional }} Estado Nacional"),
    (r"^◻ Estado Provincial", "{{ cb_origen_provincial }} Estado Provincial"),
    (r"^◻ Estado Municipal", "{{ cb_origen_municipal }} Estado Municipal"),
    (r"^◻ Donaciones", "{{ cb_origen_donaciones }} Donaciones"),
    (r"^◻ Otros$", "{{ cb_origen_otros }} Otros"),
    (
        r"^\s*◻ Reconocimiento económico",
        "{{ cb_mejora_1 }} Reconocimiento económico a talleristas",
    ),
    (
        r"^\s*◻ Incorporación de nuevos talleristas",
        "{{ cb_mejora_2 }} Incorporación de nuevos talleristas",
    ),
    (
        r"^\s*◻ Ampliación de días",
        "{{ cb_mejora_3 }} Ampliación de días y horarios de las actividades o de "
        "su carga horaria",
    ),
    (
        r"^\s*◻ Incorporación de nuevas actividades",
        "{{ cb_mejora_4 }} Incorporación de nuevas actividades",
    ),
    (
        r"^\s*◻ Mayor participación",
        "{{ cb_mejora_5 }} Mayor participación de los asistentes a las actividades",
    ),
    (
        r"^\s*◻ Calidad de los insumos",
        "{{ cb_mejora_6 }} Calidad de los insumos/materiales/equipamiento para el "
        "desarrollo de las actividades",
    ),
    (
        r"^\s*◻ Problemas con la retribución",
        "{{ cb_dif_1 }} Problemas con la retribución/facturación de los talleristas",
    ),
    (
        r"^\s*◻ Problemas con la facturación",
        "{{ cb_dif_2 }} Problemas con la facturación de los "
        "insumos/materiales/equipamiento",
    ),
    (
        r"^\s*◻ Problemas con las rendiciones",
        "{{ cb_dif_3 }} Problemas con las rendiciones de cuentas",
    ),
    (
        r"^¿Reciben demandas",
        "¿Reciben demandas para realizar otras actividades?  No {{ cb_dem_no }} "
        "Sí {{ cb_dem_si }} Cuál/es? {{ demandas_cuales }}",
    ),
    (r"^\s*☐ Presencial", "{{ cb_presencial }} Presencial"),
    (r"^\s*☐ Virtual", "{{ cb_virtual }} Virtual"),
]
# "Otras. Especificar:" y "Ninguna" aparecen dos veces (6 y 7): se distinguen por orden.
OTRAS = [
    "{{ cb_mejora_7 }} Otras. Especificar: {{ mejoras_otras }}",
    "{{ cb_dif_4 }} Otras. Especificar: {{ dificultades_otras }}",
]
NINGUNA = ["{{ cb_mejora_8 }} Ninguna", "{{ cb_dif_5 }} Ninguna"]
# Blancos sueltos (párrafo de solo puntos): 5.3.2 destino y 9 comentarios, en ese orden.
SUELTOS = ["{{ otros_recursos_destino }}", "{{ comentarios_finales }}"]
# Tablas comunes: celda vacía a la derecha de la etiqueta.
CELDAS = {
    "Nombre y Apellido": "{{ ent_nombre_apellido }}",
    "DNI del entrevistado": "{{ ent_dni }}",
    "Calle": "{{ calle }}",
    "Número": "{{ numero }}",
    "Manzana": "{{ manzana }}",
    "Lote": "{{ lote }}",
    "Barrio": "{{ barrio }}",
    "Localidad": "{{ localidad }}",
    "Municipio": "{{ municipio }}",
    "Provincia": "{{ provincia }}",
    "Firma del entrevistado": "{{ firma }}",
    "Nombre y apellido del técnico": "{{ tecnico_nombre }}",
}
COLUMNAS_ACTIVIDADES = ["tipo", "subtipo", "dias", "horario", "personas"]


def _texto(parrafo):
    return "".join(t.text or "" for t in parrafo.iter(qn("w:t")))


def _reescribir(parrafo, nuevo):
    textos = list(parrafo.iter(qn("w:t")))
    if not textos:
        return
    textos[0].text = nuevo
    textos[0].set(XML_SPACE, "preserve")
    for texto in textos[1:]:
        texto.text = ""


def _reemplazar_parrafos(doc):
    otras, ninguna, sueltos = iter(OTRAS), iter(NINGUNA), iter(SUELTOS)
    for parrafo in doc.element.iter(qn("w:p")):
        texto = _texto(parrafo)
        if not texto.strip():
            continue
        sangria = texto[: len(texto) - len(texto.lstrip())]
        for patron, nuevo in REGLAS:
            if re.search(patron, texto):
                _reescribir(parrafo, sangria + nuevo)
                break
        else:
            if re.match(r"^\s*◻ Otras", texto):
                _reescribir(parrafo, sangria + next(otras))
            elif re.match(r"^\s*◻ Ninguna", texto):
                _reescribir(parrafo, sangria + next(ninguna))
            elif PUNTOS.fullmatch(texto.strip()):
                _reescribir(parrafo, next(sueltos))


def _reemplazar_tablas(doc):
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = fila.cells
            if len(celdas) == 2 and celdas[0].text.strip() in CELDAS:
                celdas[1].paragraphs[0].add_run(CELDAS[celdas[0].text.strip()])
    # Cuadro de actividades: 6 filas de datos.
    actividades = next(t for t in doc.tables if len(t.columns) == 5)
    for i, fila in enumerate(actividades.rows[1:]):
        for j, clave in enumerate(COLUMNAS_ACTIVIDADES):
            fila.cells[j].paragraphs[0].add_run(f"{{{{ actividades[{i}].{clave} }}}}")


def _sin_fuentes(origen, destino):
    """Copia el docx sin ``word/fonts/*`` ni las referencias a esas fuentes."""
    with (
        zipfile.ZipFile(origen) as zin,
        zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for item in zin.infolist():
            if item.filename.startswith("word/fonts/"):
                continue
            datos = zin.read(item.filename)
            if item.filename == "word/fontTable.xml":
                datos = re.sub(
                    rb"<w:embed(Regular|Bold|Italic|BoldItalic)[^>]*/>", b"", datos
                )
            elif item.filename == "word/_rels/fontTable.xml.rels":
                datos = re.sub(rb"<Relationship[^>]*fonts/font[^>]*/>", b"", datos)
            elif item.filename == "word/settings.xml":
                datos = datos.replace(b"<w:embedTrueTypeFonts/>", b"").replace(
                    b"<w:saveSubsetFonts/>", b""
                )
            elif item.filename == "[Content_Types].xml":
                datos = re.sub(rb"<Default[^>]*odttf[^>]*/>", b"", datos)
            zout.writestr(item, datos)


def _sin_comentarios(doc):
    """Quita los comentarios del área (Gescom): son indicaciones para el desarrollo,
    no parte del formulario (#2651, QA del 14/10). Se borran las marcas del texto
    y la relación con comments.xml, así esa parte no se guarda."""
    cuerpo = doc.element
    for marca in ("w:commentRangeStart", "w:commentRangeEnd"):
        for nodo in list(cuerpo.iter(qn(marca))):
            nodo.getparent().remove(nodo)
    for referencia in list(cuerpo.iter(qn("w:commentReference"))):
        corrida = referencia.getparent()
        corrida.getparent().remove(corrida)
    for rid, rel in list(doc.part.rels.items()):
        if "comments" in rel.reltype.lower():
            doc.part.drop_rel(rid)


def _resumen(destino):
    with zipfile.ZipFile(destino) as plantilla:
        xml = plantilla.read("word/document.xml").decode("utf8")
        plantilla_nombres = plantilla.namelist()
    tags = re.findall(r"\{\{ ([a-z_0-9\[\]\.]+) \}\}", xml)
    resto = re.sub(r"<[^>]+>", "", xml)
    comentarios = [n for n in plantilla_nombres if "comment" in n.lower()]
    if comentarios or "comment" in xml.lower():
        raise SystemExit(f"La plantilla conserva comentarios: {comentarios}")
    for texto_de_trabajo in ("Gescom", "GESCOM", "precargado en la app", "Comentario"):
        if texto_de_trabajo in resto:
            raise SystemExit(
                f"La plantilla conserva texto de trabajo: {texto_de_trabajo!r}"
            )
    print(f"tags: {len(tags)} ({len(set(tags))} distintos)")
    print(
        "puntos sin reemplazar:",
        len(PUNTOS.findall(resto)),
        "| casillas sin reemplazar:",
        resto.count("◻") + resto.count("☐"),
    )


def main(origen, destino):
    doc = Document(origen)
    _reemplazar_parrafos(doc)
    _reemplazar_tablas(doc)
    _sin_comentarios(doc)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as con_fuentes:
        doc.save(con_fuentes.name)
        _sin_fuentes(con_fuentes.name, destino)
    _resumen(destino)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
