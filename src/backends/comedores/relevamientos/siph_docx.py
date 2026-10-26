"""Exportación del Seguimiento SIPH al docx del área (#2651).

La plantilla ``templates/relevamientos/docx/seguimiento_siph.docx`` es el
formulario que adjuntó el área ("IV-Formulario Seguimiento de SIPH-Línea
Tradicional y Módulos secos"), con sus blancos, casillas y celdas reemplazados
por tags de docxtpl. Las respuestas salen de ``SeguimientoPnud.datos`` con las
claves del formulario de la app (``pnud_formularios.json``, ``siph``).
"""

import io
import subprocess
import tempfile
from pathlib import Path
from urllib.parse import unquote, urlparse

from django.core.files.storage import default_storage
from django.utils import timezone
from docx import Document
from docx.shared import Cm

# El paquete docx_service se reemplaza a si mismo por `impl` al importarse: pylint
# no ve sus nombres, asi que se importa el modulo real (es el mismo objeto).
from admisiones.services.docx_service.impl import DocumentTemplateService
from relevamientos.pnud_formularios import definiciones

PLANTILLA = "seguimiento_siph.docx"
FILAS_ACTIVIDADES = 6
# "☐" (U+2610) no está en las fuentes del servidor y LibreOffice lo dibuja roto;
# "□" y "☒" salen de la misma fuente (OpenSymbol).
MARCADA, VACIA = "☒", "□"


def _texto(valor):
    return "" if valor in (None, "") else str(valor)


def _lista(valor):
    if isinstance(valor, list):
        return [str(v) for v in valor]
    return [] if valor in (None, "") else [str(valor)]


def _marca(condicion):
    return MARCADA if condicion else VACIA


def _opciones_siph(nombre):
    """Valores de las opciones de un campo del SIPH, en el orden del papel."""
    for seccion in (definiciones().get("siph") or {}).get("sections", []):
        for campo in seccion.get("fields", []):
            if campo.get("name") == nombre:
                return [opcion["value"] for opcion in campo.get("options", [])]
    return []


def _marcas(prefijo, nombre, elegidas):
    """``cb_<prefijo>_1..n``: una casilla por opción del papel, en su orden."""
    return {
        f"cb_{prefijo}_{indice}": _marca(valor in elegidas)
        for indice, valor in enumerate(_opciones_siph(nombre), start=1)
    }


def _actividades(filas):
    """Las 6 filas del cuadro 4; las que faltan van vacías."""
    salida = []
    for fila in filas[:FILAS_ACTIVIDADES]:
        fila = fila if isinstance(fila, dict) else {}
        salida.append(
            {
                "tipo": _texto(fila.get("tipo_actividad")),
                "subtipo": _texto(fila.get("subtipo_actividad")),
                "dias": _texto(fila.get("dias_funcionamiento")),
                "horario": _texto(fila.get("horario_funcionamiento")),
                "personas": _texto(fila.get("personas_observadas")),
            }
        )
    vacia = {"tipo": "", "subtipo": "", "dias": "", "horario": "", "personas": ""}
    return salida + [dict(vacia) for _ in range(FILAS_ACTIVIDADES - len(salida))]


def _fecha(valor):
    return timezone.localtime(valor).strftime("%d/%m/%Y %H:%M") if valor else ""


def _nombre(usuario):
    if usuario is None:
        return ""
    return usuario.get_full_name() or usuario.get_username()


def contexto_docx_siph(seguimiento):  # pylint: disable=too-many-locals
    datos = seguimiento.datos or {}
    comedor = seguimiento.comedor
    funcionamiento = _texto(datos.get("funcionamiento"))
    if funcionamiento == "Cerrado":
        # Con "Cerrado" la app deriva al formulario de excepción: ese es el motivo.
        motivo = " — ".join(
            parte
            for parte in (datos.get("excepcion_motivo"), datos.get("excepcion_detalle"))
            if parte
        )
    else:
        motivo = _texto(datos.get("funcionamiento_motivo"))
    actividad = _texto(datos.get("actividad_visitada"))
    if actividad == "Otra":
        actividad = _texto(datos.get("actividad_visitada_otra"))
    organizacion = getattr(comedor, "organizacion", None)
    contexto = {
        "organizacion": _texto(getattr(organizacion, "nombre", "")),
        "codigo_proyecto": _texto(
            datos.get("codigo_proyecto") or comedor.codigo_de_proyecto
        ),
        "comedor_nombre": _texto(comedor.nombre),
        "codigo_externo": _texto(datos.get("codigo_externo") or comedor.id_externo),
        "id_sisoc": _texto(comedor.id),
        "id_relevamiento": _texto(seguimiento.id),
        "fecha_hora": _fecha(seguimiento.fecha_hora),
        "x_abierto_con": "X" if funcionamiento == "Abierto, con actividad" else "",
        "x_abierto_sin": "X" if funcionamiento == "Abierto, sin actividad" else "",
        "x_cerrado": "X" if funcionamiento == "Cerrado" else "",
        "motivo": motivo,
        "ent_nombre_apellido": _texto(datos.get("ent_nombre_apellido")),
        "ent_dni": _texto(datos.get("ent_dni")),
        "cb_actividad_si": _marca(datos.get("actividad_en_espacio") is True),
        "cb_actividad_no": _marca(datos.get("actividad_en_espacio") is False),
        "actividad_en_espacio_motivo": _texto(datos.get("actividad_en_espacio_motivo")),
        "actividad_nombre": actividad,
        "actividades": _actividades(datos.get("actividades") or []),
        "cb_fin_si": _marca(datos.get("financiamiento_planificado") is True),
        "cb_fin_no": _marca(datos.get("financiamiento_planificado") is False),
        "financiamiento_talleristas_cantidad": _texto(
            datos.get("financiamiento_talleristas_cantidad")
        ),
        "financiamiento_materiales_detalle": _texto(
            datos.get("financiamiento_materiales_detalle")
        ),
        "financiamiento_adecuacion_detalle": _texto(
            datos.get("financiamiento_adecuacion_detalle")
        ),
        "financiamiento_justificacion": _texto(
            datos.get("financiamiento_justificacion")
        ),
        "cb_otros_si": _marca(datos.get("otros_recursos") is True),
        "cb_otros_no": _marca(datos.get("otros_recursos") is False),
        "otros_recursos_destino": _texto(datos.get("otros_recursos_destino")),
        "mejoras_otras": _texto(datos.get("mejoras_otras")),
        "dificultades_otras": _texto(datos.get("dificultades_otras")),
        "cb_dem_si": _marca(datos.get("demandas_otras_actividades") is True),
        "cb_dem_no": _marca(datos.get("demandas_otras_actividades") is False),
        "demandas_cuales": _texto(datos.get("demandas_cuales")),
        "comentarios_finales": _texto(datos.get("comentarios_finales")),
        "cb_presencial": _marca(datos.get("modalidad_entrevista") == "Presencial"),
        "cb_virtual": _marca(datos.get("modalidad_entrevista") == "Virtual"),
        # La firma va como imagen en su celda (ver _con_firma), no como texto.
        "firma": "",
        "tecnico_nombre": _nombre(seguimiento.tecnico),
    }
    for clave in (
        "calle",
        "numero",
        "manzana",
        "lote",
        "barrio",
        "localidad",
        "municipio",
        "provincia",
    ):
        contexto[clave] = _texto(datos.get(clave))
    componentes = _lista(datos.get("financiamiento_componentes"))
    contexto.update(
        {
            "cb_comp_talleristas": _marca("Retribución talleristas" in componentes),
            "cb_comp_materiales": _marca(
                "Materiales/insumos/equipamiento" in componentes
            ),
            "cb_comp_adecuacion": _marca(
                "Adecuación mínima de servicios básicos" in componentes
            ),
        }
    )
    origen = _lista(datos.get("otros_recursos_origen"))
    contexto.update(
        {
            "cb_origen_nacional": _marca("Estado Nacional" in origen),
            "cb_origen_provincial": _marca("Estado Provincial" in origen),
            "cb_origen_municipal": _marca("Estado Municipal" in origen),
            "cb_origen_donaciones": _marca("Donaciones" in origen),
            "cb_origen_otros": _marca("Otros" in origen),
        }
    )
    contexto.update(_marcas("mejora", "mejoras", _lista(datos.get("mejoras"))))
    contexto.update(_marcas("dif", "dificultades", _lista(datos.get("dificultades"))))
    return contexto


def _firma(seguimiento):
    """Imagen de la firma del entrevistado, o ``None``.

    La app la sube con /firma/ y en ``datos`` queda su URL, que la API ya validó
    dentro de ``firmas/<comedor>/`` del storage. Se lee del storage, nunca por red.
    """
    url = (seguimiento.datos or {}).get("firma_entrevistado")
    if not isinstance(url, str) or not url:
        return None
    ruta = unquote(urlparse(url).path)
    prefijo = urlparse(default_storage.url("")).path
    if not ruta.startswith(prefijo):
        return None
    nombre = ruta[len(prefijo) :]
    if not nombre.startswith(f"firmas/{seguimiento.comedor_id}/") or ".." in nombre:
        return None
    if not default_storage.exists(nombre):
        return None
    with default_storage.open(nombre, "rb") as archivo:
        return archivo.read()


def _con_firma(buffer, firma):
    """Pone la firma en la celda a la derecha de "Firma del entrevistado"."""
    documento = Document(buffer)
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = fila.cells
            if len(celdas) == 2 and celdas[0].text.strip() == "Firma del entrevistado":
                parrafo = celdas[1].paragraphs[0]
                parrafo.add_run().add_picture(io.BytesIO(firma), height=Cm(1.6))
    salida = io.BytesIO()
    documento.save(salida)
    salida.seek(0)
    return salida


def docx_siph(seguimiento):
    """Docx del SIPH como ``BytesIO`` listo para descargar."""
    buffer = DocumentTemplateService.generar_docx(
        PLANTILLA, contexto_docx_siph(seguimiento), app_name="relevamientos"
    )
    firma = _firma(seguimiento)
    if firma:
        try:
            return _con_firma(buffer, firma)
        except Exception:  # pylint: disable=broad-exception-caught
            # Una imagen ilegible no impide descargar el formulario.
            buffer.seek(0)
    return buffer


def pdf_siph(seguimiento):
    """PDF del SIPH (QA del 14/10): la plantilla completada, convertida con LibreOffice.

    Es el que descargan SISOC y Gestionar: los dos piden este mismo archivo.
    Misma conversión que la certificación de prestaciones y los informes de
    admisiones (perfil de LibreOffice propio en un directorio temporal).
    """
    with tempfile.TemporaryDirectory(prefix="seguimiento-siph-") as temp_dir:
        temp_path = Path(temp_dir)
        docx_path = temp_path / "seguimiento-siph.docx"
        pdf_path = temp_path / "seguimiento-siph.pdf"
        docx_path.write_bytes(docx_siph(seguimiento).getvalue())
        result = subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--nologo",
                "--nodefault",
                "--nofirststartwizard",
                f"-env:UserInstallation={(temp_path / 'perfil').resolve().as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temp_path),
                str(docx_path),
            ],
            capture_output=True,
            check=False,
            timeout=120,
        )
        if result.returncode != 0 or not pdf_path.exists():
            detalle = result.stderr.decode("utf-8", errors="replace").strip()
            raise RuntimeError(detalle or "No se pudo generar el PDF del SIPH.")
        return pdf_path.read_bytes()
