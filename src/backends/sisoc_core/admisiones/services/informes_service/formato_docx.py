"""Formato Word del Informe Técnico generado desde el Gestor de templates.

Replica el membrete del modelo institucional del Programa Alimentar Comunidad
y reúne las constantes de formato que usa el generador del informe.
"""

import logging
import os
import re

from django.apps import apps
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum import text as docx_text
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

logger = logging.getLogger("admisiones.services.informes")

PARRAFO_VACIO_EDITOR = re.compile(r"<p(\s[^>]*)?>\s*<br\s*/?>\s*</p>", re.IGNORECASE)

# Elementos que el esquema de w:tblPr ubica después de w:tblBorders.
SUCESORES_BORDES_TABLA = (
    "w:shd",
    "w:tblLayout",
    "w:tblCellMar",
    "w:tblLook",
    "w:tblCaption",
    "w:tblDescription",
)

# Elementos que el esquema de w:pPr ubica después de w:pBdr.
SUCESORES_BORDES_PARRAFO = (
    "w:shd",
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:jc",
    "w:outlineLvl",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)

# python-docx arma sus enums en tiempo de ejecución y pylint no los ve.
# pylint: disable=no-member
ALINEACIONES_HTML = {
    "left": docx_text.WD_ALIGN_PARAGRAPH.LEFT,
    "center": docx_text.WD_ALIGN_PARAGRAPH.CENTER,
    "right": docx_text.WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": docx_text.WD_ALIGN_PARAGRAPH.JUSTIFY,
}
CENTRADO_VERTICAL = WD_CELL_VERTICAL_ALIGNMENT.CENTER
TABULACION_CENTRADA = docx_text.WD_TAB_ALIGNMENT.CENTER
# pylint: enable=no-member


class MembreteInformeTecnico:
    """Encabezado y pie del modelo institucional del Informe Técnico.

    El modelo arma el encabezado con una tabla invisible; acá se replica con
    párrafos y tabulaciones porque GDE no admite tablas en el documento.
    """

    ARCHIVO_ESCUDO = "escudo_informe_tecnico.jpg"
    RENGLONES_PROGRAMA = ("PROGRAMA ALIMENTAR", "COMUNIDAD")
    TEXTO_SECRETARIA = "Secretaría de Niñez Adolescencia y Familia"
    TEXTO_MINISTERIO = "Ministerio de Capital Humano"
    FUENTE = "Times New Roman"
    COLOR_TEXTO = RGBColor(0x17, 0x37, 0x5E)
    COLOR_LINEA = "002060"
    ANCHO_ESCUDO_PT = 24
    TAMANIO_PROGRAMA_PT = 10.5
    TAMANIO_SECRETARIA_PT = 8.5
    TAMANIO_MINISTERIO_PT = 10.5
    # Ejes centrados del bloque del programa y del bloque del ministerio.
    EJE_PROGRAMA_MM = 22
    EJE_MINISTERIO_MM = 123

    @classmethod
    def aplicar(cls, documento):
        seccion = documento.sections[0]
        cls._armar_encabezado(seccion.header)
        cls._armar_pie(seccion.footer)

    @classmethod
    def _armar_encabezado(cls, encabezado):
        # El primer párrafo se conserva: la vista previa le agrega la marca de agua.
        parrafo_escudo = encabezado.paragraphs[0]
        parrafo_secretaria = encabezado.add_paragraph()
        parrafo_ministerio = encabezado.add_paragraph()
        for parrafo in (parrafo_escudo, parrafo_secretaria, parrafo_ministerio):
            cls._preparar_renglon(parrafo)

        ruta_escudo = cls._ruta_escudo()
        if os.path.exists(ruta_escudo):
            parrafo_escudo.add_run("\t\t")
            parrafo_escudo.add_run().add_picture(
                ruta_escudo, width=Pt(cls.ANCHO_ESCUDO_PT)
            )
        else:
            logger.warning("No se encontró el escudo del membrete: %s", ruta_escudo)

        programa_superior, programa_inferior = cls.RENGLONES_PROGRAMA
        cls._agregar_texto(
            parrafo_secretaria, f"\t{programa_superior}", cls.TAMANIO_PROGRAMA_PT
        )
        cls._agregar_texto(
            parrafo_secretaria, f"\t{cls.TEXTO_SECRETARIA}", cls.TAMANIO_SECRETARIA_PT
        )
        cls._agregar_texto(
            parrafo_ministerio, f"\t{programa_inferior}", cls.TAMANIO_PROGRAMA_PT
        )
        cls._agregar_texto(
            parrafo_ministerio, f"\t{cls.TEXTO_MINISTERIO}", cls.TAMANIO_MINISTERIO_PT
        )
        cls._aplicar_linea_inferior(parrafo_ministerio)

    @classmethod
    def _armar_pie(cls, pie):
        parrafo = pie.paragraphs[0]
        parrafo.alignment = ALINEACIONES_HTML["right"]
        campo = OxmlElement("w:fldSimple")
        campo.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        texto = OxmlElement("w:t")
        texto.text = "1"
        run.append(texto)
        campo.append(run)
        parrafo._p.append(campo)  # pylint: disable=protected-access

    @classmethod
    def _preparar_renglon(cls, parrafo):
        # Sin el estilo "Header", que trae sus propias tabulaciones.
        parrafo.style = "Normal"
        parrafo.alignment = ALINEACIONES_HTML["left"]
        formato = parrafo.paragraph_format
        formato.space_before = 0
        formato.space_after = 0
        formato.line_spacing = 1
        for eje_mm in (cls.EJE_PROGRAMA_MM, cls.EJE_MINISTERIO_MM):
            formato.tab_stops.add_tab_stop(Mm(eje_mm), TABULACION_CENTRADA)

    @classmethod
    def _agregar_texto(cls, parrafo, texto, tamanio_pt):
        run = parrafo.add_run(texto)
        run.font.name = cls.FUENTE
        run.font.size = Pt(tamanio_pt)
        run.font.color.rgb = cls.COLOR_TEXTO

    @classmethod
    def _aplicar_linea_inferior(cls, parrafo):
        propiedades = parrafo._p.get_or_add_pPr()  # pylint: disable=protected-access
        bordes = OxmlElement("w:pBdr")
        borde = OxmlElement("w:bottom")
        borde.set(qn("w:val"), "single")
        borde.set(qn("w:sz"), "12")
        borde.set(qn("w:space"), "8")
        borde.set(qn("w:color"), cls.COLOR_LINEA)
        bordes.append(borde)
        propiedades.insert_element_before(bordes, *SUCESORES_BORDES_PARRAFO)

    @classmethod
    def _ruta_escudo(cls):
        return os.path.join(
            apps.get_app_config("admisiones").path,
            "templates",
            "admisiones",
            "docx",
            cls.ARCHIVO_ESCUDO,
        )
